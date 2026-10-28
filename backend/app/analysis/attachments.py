"""Attachment validation, safe text extraction and document intelligence.

Attachments are untrusted: they are never executed, links inside them are never
followed, and their text is passed to the LLM only as delimited data.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any
from uuid import UUID

from sqlalchemy.dialects.postgresql import insert as pg_insert

from app.analysis.deadlines import extract_deadlines
from app.analysis.phishing import DANGEROUS_EXTENSIONS, MACRO_EXTENSIONS
from app.config import get_settings
from app.db.models import EmailAttachment, EmailMessage, SearchDocument, UserSettings
from app.db.session import system_session
from app.errors import NotFoundError
from app.gmail.client import GmailClient
from app.gmail.parser import b64url_decode
from app.llm.base import LLMError, LLMNotConfiguredError
from app.llm.prompts import ATTACHMENT_SYSTEM
from app.llm.schemas import AttachmentAnalysisOutput
from app.llm.service import LLMService, get_llm_service
from app.logging import get_logger

log = get_logger(__name__)

SUPPORTED = {
    "application/pdf": "pdf",
    "text/plain": "text",
    "text/csv": "csv",
    "application/csv": "csv",
    "application/vnd.ms-excel": "csv",  # many mailers label CSV like this; content is sniffed below
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}
EXT_KIND = {".pdf": "pdf", ".txt": "text", ".csv": "csv", ".docx": "docx", ".md": "text", ".log": "text"}

_AMOUNT = re.compile(
    r"(?P<cur>₹|\$|€|£|INR|USD|EUR|GBP|Rs\.?)\s?(?P<val>\d{1,3}(?:[,\s]\d{2,3})*(?:\.\d{1,2})?|\d+(?:\.\d{1,2})?)"
    r"|(?P<val2>\d{1,3}(?:,\d{2,3})+(?:\.\d{1,2})?)\s?(?P<cur2>INR|USD|EUR|GBP|rupees|dollars)",
    re.I,
)
_INVOICE = re.compile(r"\b(?:invoice|inv|bill)\s*(?:no\.?|number|#|:)\s*[:#]?\s*(?P<id>[A-Z0-9][A-Z0-9\-/]{2,30})", re.I)
_QUOTE = re.compile(r"\b(?:quotation|quote|qtn|proposal)\s*(?:no\.?|number|#|ref\.?)?\s*[:#]?\s*(?P<id>[A-Z]{1,5}-?\d[\w\-/]{1,25})", re.I)
_INJECTION = re.compile(
    r"\b(ignore (?:all |any )?(?:previous|prior|above) (?:instructions|rules)|you are now|system prompt|"
    r"disregard (?:your|the) (?:instructions|rules)|forward (?:all|every) (?:emails?|messages?))\b",
    re.I,
)


@dataclass(slots=True)
class ValidationResult:
    ok: bool
    kind: str | None
    status: str
    reason: str | None = None


class ExtractionError(Exception):
    pass


def _ext(filename: str) -> str:
    return "." + filename.lower().rsplit(".", 1)[-1] if "." in filename else ""


def validate_attachment(filename: str, declared_mime: str, size: int, max_bytes: int) -> ValidationResult:
    ext = _ext(filename)
    if ext in DANGEROUS_EXTENSIONS or ext in MACRO_EXTENSIONS:
        return ValidationResult(False, None, "blocked", "Executable, script or macro attachments are never processed")
    hard = get_settings().attachment_hard_max_bytes
    if size > min(max_bytes, hard):
        return ValidationResult(False, None, "too_large", f"Attachment exceeds the {min(max_bytes, hard)} byte limit")
    kind = SUPPORTED.get(declared_mime.lower()) or EXT_KIND.get(ext)
    if kind is None:
        return ValidationResult(False, None, "unsupported", f"Unsupported type {declared_mime}")
    return ValidationResult(True, kind, "pending")


def sniff_kind(data: bytes, expected: str) -> str:
    """Confirm content matches the expected kind using magic bytes; raise on mismatch."""
    if expected == "pdf":
        if not data[:1024].lstrip().startswith(b"%PDF-"):
            raise ExtractionError("File is not a valid PDF")
        return "pdf"
    if expected == "docx":
        if not data.startswith(b"PK\x03\x04"):
            raise ExtractionError("File is not a valid DOCX package")
        return "docx"
    if b"\x00" in data[:4096]:
        raise ExtractionError("Binary content where text was expected")
    return expected


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _ocr_pdf(data: bytes, max_pages: int) -> str:
    settings = get_settings()
    if not settings.ocr_enabled:
        raise ExtractionError("Scanned PDF without a text layer; OCR is not enabled (set OCR_ENABLED=true)")
    try:
        import pypdfium2 as pdfium  # type: ignore[import-not-found]
        import pytesseract  # type: ignore[import-not-found]
    except ImportError as exc:
        raise ExtractionError("OCR dependencies are not installed (pip install '.[ocr]')") from exc
    if settings.tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
    pdf = pdfium.PdfDocument(data)
    texts = []
    for i in range(min(len(pdf), max_pages)):
        image = pdf[i].render(scale=2).to_pil()
        texts.append(pytesseract.image_to_string(image))
    return "[OCR text — may contain recognition errors]\n" + "\n".join(texts)


def extract_text(data: bytes, kind: str) -> tuple[str, dict[str, Any]]:
    settings = get_settings()
    limit = settings.attachment_max_extract_chars
    meta: dict[str, Any] = {}
    if kind == "pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ExtractionError("Encrypted PDF cannot be processed")
            pages = reader.pages
            meta["pages"] = len(pages)
            parts = []
            for page in list(pages)[: settings.pdf_max_pages]:
                parts.append(page.extract_text() or "")
                if sum(map(len, parts)) > limit:
                    break
            text = "\n".join(parts).strip()
        except PdfReadError as exc:
            raise ExtractionError(f"Malformed PDF: {exc}") from exc
        if len(text) < 20:
            text = _ocr_pdf(data, settings.pdf_max_pages)
            meta["ocr"] = True
    elif kind == "docx":
        import docx  # python-docx

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:  # malformed zip/xml
            raise ExtractionError("Malformed DOCX document") from exc
        lines = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables[:20]:
            for row in table.rows[:200]:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(lines)
    elif kind == "csv":
        raw = _decode_text(data)
        try:
            dialect = csv.Sniffer().sniff(raw[:4096]) if raw.strip() else csv.excel
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(raw), dialect))
        meta["rows"] = len(rows)
        meta["columns"] = len(rows[0]) if rows else 0
        text = "\n".join(" | ".join(cell[:200] for cell in row[:30]) for row in rows[:500])
    else:
        text = _decode_text(data)
    return text[:limit], meta


def deterministic_extraction(text: str, now: datetime, tz: str) -> dict[str, Any]:
    amounts = []
    for m in _AMOUNT.finditer(text[:50000]):
        cur = (m.group("cur") or m.group("cur2") or "").upper().replace("RS.", "INR").replace("RS", "INR")
        cur = {"₹": "INR", "$": "USD", "€": "EUR", "£": "GBP", "RUPEES": "INR", "DOLLARS": "USD"}.get(cur, cur)
        val = (m.group("val") or m.group("val2") or "").strip()
        ctx = text[max(0, m.start() - 60): m.end() + 20].replace("\n", " ")
        amounts.append({"value": val, "currency": cur or None, "context": ctx.strip()[:200]})
    deadlines = [
        {"description": c.text, "iso": c.due_at.isoformat(), "evidence": c.text, "date_type": c.date_type}
        for c in extract_deadlines(text, now, tz)[:10]
    ]
    return {
        "amounts": amounts[:25],
        "invoice_numbers": sorted({m.group("id") for m in _INVOICE.finditer(text)})[:10],
        "quotation_refs": sorted({m.group("id") for m in _QUOTE.finditer(text)})[:10],
        "deadlines": deadlines,
        "prompt_injection_suspected": bool(_INJECTION.search(text)),
    }


async def _fetch_bytes(gmail: GmailClient, msg: EmailMessage, att: EmailAttachment) -> bytes:
    att_id = (att.extraction or {}).get("gmail_attachment_id")
    if not att_id:
        # Gmail attachment IDs are not stable across fetches: re-read the message to get a fresh one.
        raw = await gmail.get_message(msg.gmail_message_id, "full")
        stack = [raw.get("payload") or {}]
        while stack:
            part = stack.pop()
            stack.extend(part.get("parts") or [])
            if str(part.get("partId")) == att.part_id:
                body = part.get("body") or {}
                if body.get("data"):
                    return b64url_decode(body["data"])
                att_id = body.get("attachmentId")
                break
    if not att_id:
        raise ExtractionError("Attachment content could not be located in Gmail")
    try:
        data = await gmail.get_attachment(msg.gmail_message_id, att_id)
    except Exception:
        raw = await gmail.get_message(msg.gmail_message_id, "full")
        stack = [raw.get("payload") or {}]
        fresh = None
        while stack:
            part = stack.pop()
            stack.extend(part.get("parts") or [])
            if str(part.get("partId")) == att.part_id:
                fresh = (part.get("body") or {}).get("attachmentId")
        if not fresh:
            raise
        data = await gmail.get_attachment(msg.gmail_message_id, fresh)
    return b64url_decode(data.get("data", ""))


async def process_attachment(
    user_id: UUID, attachment_id: UUID, gmail: GmailClient | None = None, llm: LLMService | None = None
) -> EmailAttachment:
    llm = llm or get_llm_service()
    async with system_session() as session:
        att = await session.get(EmailAttachment, attachment_id)
        if att is None or att.user_id != user_id:
            raise NotFoundError("Attachment not found")
        msg = await session.get(EmailMessage, att.message_id)
        settings = await session.get(UserSettings, user_id) or UserSettings(user_id=user_id)
        assert msg is not None

    async def _save(**values: Any) -> EmailAttachment:
        async with system_session() as s:
            row = await s.get(EmailAttachment, attachment_id)
            assert row is not None
            for k, v in values.items():
                setattr(row, k, v)
            await s.flush()
            return row

    validation = validate_attachment(att.filename, att.mime_type, att.size_bytes, settings.attachment_max_bytes)
    if not validation.ok:
        flags = [{"code": "dangerous_attachment", "detail": validation.reason}] if validation.status == "blocked" else []
        return await _save(extraction_status=validation.status, extraction_error=validation.reason, risk_flags=flags)

    owns = gmail is None
    if gmail is None:
        from app.gmail.oauth import gmail_client_for

        gmail = await gmail_client_for(msg.account_id)
    try:
        data = await _fetch_bytes(gmail, msg, att)
        if len(data) > min(settings.attachment_max_bytes, get_settings().attachment_hard_max_bytes):
            return await _save(extraction_status="too_large", extraction_error="Attachment exceeds size limit")
        kind = sniff_kind(data, validation.kind or "text")
        text, meta = extract_text(data, kind)
    except ExtractionError as exc:
        return await _save(extraction_status="failed", extraction_error=str(exc)[:500])
    finally:
        if owns:
            await gmail.aclose()

    digest = hashlib.sha256(data).hexdigest()
    now = datetime.now(UTC)
    deterministic = deterministic_extraction(text, now, settings.timezone)
    extraction: dict[str, Any] = {"deterministic": deterministic, "meta": meta, "source": "rules"}
    risk_flags: list[dict[str, str]] = []
    if deterministic["prompt_injection_suspected"]:
        risk_flags.append({"code": "prompt_injection", "detail": "Document contains instructions aimed at AI tools"})

    if text.strip():
        try:
            out, meta_llm = await llm.generate(
                task="attachment_analysis", system_prompt=ATTACHMENT_SYSTEM,
                untrusted={"document": text, "email_subject": msg.subject or ""},
                trusted_context=f"Filename: {att.filename!r} (untrusted). Current date: {now.date()} "
                                f"(user timezone {settings.timezone}).",
                output_model=AttachmentAnalysisOutput, user_id=user_id,
            )
            extraction["llm"] = out.model_dump()
            extraction["model"] = meta_llm.model
            extraction["source"] = "hybrid"
            if out.prompt_injection_detected and not risk_flags:
                risk_flags.append({"code": "prompt_injection", "detail": "Document contains instructions aimed at AI tools"})
        except LLMNotConfiguredError:
            pass
        except LLMError as exc:
            extraction["llm_error"] = exc.code

    storage_path = None
    if settings.store_attachments:
        from app.services.storage import safe_object_name, upload_object

        storage_path = f"{user_id}/{attachment_id}/{safe_object_name(att.filename)}"
        try:
            await upload_object(storage_path, data, att.mime_type)
        except Exception as exc:  # storage optional; analysis still useful
            log.warning("attachment_store_failed", error=type(exc).__name__)
            storage_path = None

    async with system_session() as session:
        await session.execute(
            pg_insert(SearchDocument).values(
                user_id=user_id, message_id=att.message_id, attachment_id=attachment_id,
                title=att.filename, content=text[:150_000],
            ).on_conflict_do_update(
                index_elements=["message_id", "attachment_id"], set_={"content": text[:150_000], "title": att.filename}
            )
        )
    return await _save(
        extraction_status="done", extraction_error=None, extracted_text=text, extraction=extraction,
        risk_flags=risk_flags, sha256=digest, storage_path=storage_path, size_bytes=len(data),
    )
